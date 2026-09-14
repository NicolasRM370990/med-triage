"""Gera a apresentação PPT do projeto Med-Triage."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
TEAL = RGBColor(0x1A, 0xA6, 0xA6)
GOLD = RGBColor(0xE8, 0xB8, 0x4A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFFWHITE = RGBColor(0xF4, 0xF7, 0xFA)
SLATE = RGBColor(0x3A, 0x4A, 0x5C)
MUTED = RGBColor(0x6B, 0x7C, 0x8D)
CARD = RGBColor(0x12, 0x2E, 0x52)
GREEN = RGBColor(0x2E, 0xA8, 0x6A)

W = Inches(13.333)
H = Inches(7.5)
OUT = Path(__file__).resolve().parent / "Med-Triage_Analise_STAR.pptx"


def set_run(run, size, bold=False, color=WHITE, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_text_box(slide, l, t, w, h, text, size, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return tf


def add_para(tf, text, size, bold=False, color=SLATE, space_before=6, space_after=4):
    p = tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return p


def rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(1, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def card(slide, l, t, w, h, fill=WHITE):
    return rect(slide, l, t, w, h, fill)


def accent_bar(slide, l, t, w=Inches(0.08), h=Inches(0.9)):
    return rect(slide, l, t, w, h, TEAL)


def footer(slide, page, total=12):
    rect(slide, Inches(0), Inches(7.22), W, Inches(0.28), NAVY)
    add_text_box(
        slide,
        Inches(0.4),
        Inches(7.22),
        Inches(9),
        Inches(0.28),
        "Med-Triage  ·  Classificação de abstracts médicos  ·  Tech Challenge",
        10,
        False,
        RGBColor(0xB8, 0xC5, 0xD6),
    )
    add_text_box(
        slide,
        Inches(11.4),
        Inches(7.22),
        Inches(1.5),
        Inches(0.28),
        f"{page:02d}  /  {total:02d}",
        10,
        False,
        GOLD,
        PP_ALIGN.RIGHT,
    )


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def header_slide(prs, kicker, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, OFFWHITE)
    rect(slide, 0, 0, W, Inches(1.35), NAVY)
    rect(slide, 0, Inches(1.35), W, Inches(0.08), TEAL)
    add_text_box(slide, Inches(0.5), Inches(0.18), Inches(8), Inches(0.3), kicker.upper(), 12, True, TEAL)
    add_text_box(slide, Inches(0.5), Inches(0.48), Inches(12), Inches(0.7), title, 28, True, WHITE)
    if subtitle:
        add_text_box(slide, Inches(0.5), Inches(1.52), Inches(12), Inches(0.4), subtitle, 14, False, MUTED)
    return slide


def bullet_card(slide, l, t, w, h, title, items, title_color=NAVY):
    card(slide, l, t, w, h, WHITE)
    accent_bar(slide, l, t, Inches(0.08), h)
    tf = add_text_box(
        slide,
        l + Inches(0.28),
        t + Inches(0.12),
        w - Inches(0.4),
        Inches(0.4),
        title,
        15,
        True,
        title_color,
    )
    for item in items:
        add_para(tf, "•  " + item, 13, False, SLATE, 8, 2)


def cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, NAVY)
    rect(slide, 0, 0, Inches(0.18), H, TEAL)
    rect(slide, Inches(0.18), Inches(6.55), W, Inches(0.95), CARD)
    add_text_box(
        slide,
        Inches(0.7),
        Inches(1.35),
        Inches(12),
        Inches(0.35),
        "TECH CHALLENGE  ·  NLP + MLOps",
        14,
        True,
        TEAL,
    )
    add_text_box(slide, Inches(0.7), Inches(1.85), Inches(12), Inches(1.4), "Med-Triage", 54, True, WHITE)
    add_text_box(
        slide,
        Inches(0.7),
        Inches(3.35),
        Inches(11.5),
        Inches(1.2),
        "API de classificação de abstracts médicos em cinco categorias clínicas,\n"
        "com pipeline TF-IDF + regressão logística, Docker e CI/CD.",
        18,
        False,
        RGBColor(0xC9, 0xD6, 0xE4),
    )
    add_text_box(
        slide,
        Inches(0.7),
        Inches(5.0),
        Inches(11),
        Inches(0.4),
        "Análise de projeto  ·  Apresentação no modelo STAR  ·  até 5 minutos",
        16,
        False,
        GOLD,
    )
    add_text_box(
        slide,
        Inches(0.7),
        Inches(6.7),
        Inches(12),
        Inches(0.5),
        "Medical Abstracts TC Corpus  ·  FastAPI  ·  scikit-learn  ·  Prometheus  ·  GitHub Actions",
        13,
        False,
        RGBColor(0xB8, 0xC5, 0xD6),
    )
    notes(
        slide,
        "Abertura (0:00–0:20). Diga o nome do projeto e a proposta em uma frase: "
        "classificar abstracts médicos e servir o modelo via API de apoio à triagem.",
    )


def agenda(prs):
    slide = header_slide(
        prs,
        "Roteiro",
        "Apresentação no modelo STAR",
        "Até 5 minutos  ·  Situation, Task, Action, Result",
    )
    blocks = [
        ("S", "Situation", "O volume de literatura biomédica e o corpus de 5 classes."),
        ("T", "Task", "Classificar, avaliar e servir um MVP operacional."),
        ("A", "Action", "Arquitetura, modelo, API, testes e esteira."),
        ("R", "Result", "Métricas reais, o que funciona e o próximo salto."),
    ]
    x = Inches(0.5)
    for letter, name, desc in blocks:
        card(slide, x, Inches(2.15), Inches(2.9), Inches(4.4), WHITE)
        rect(slide, x, Inches(2.15), Inches(2.9), Inches(1.15), NAVY)
        add_text_box(slide, x, Inches(2.25), Inches(2.9), Inches(0.9), letter, 36, True, TEAL, PP_ALIGN.CENTER)
        add_text_box(
            slide,
            x + Inches(0.15),
            Inches(3.5),
            Inches(2.6),
            Inches(0.5),
            name,
            18,
            True,
            NAVY,
            PP_ALIGN.CENTER,
        )
        add_text_box(
            slide,
            x + Inches(0.2),
            Inches(4.15),
            Inches(2.5),
            Inches(1.8),
            desc,
            14,
            False,
            SLATE,
            PP_ALIGN.CENTER,
        )
        x += Inches(3.15)
    footer(slide, 2)
    notes(slide, "Agenda (0:20–0:35). Avise STAR e que o foco é produto + operação, não só o modelo.")


def situation(prs):
    slide = header_slide(
        prs,
        "Situation",
        "Por que classificar abstracts médicos?",
        "Contexto clínico e de produto — apoio à triagem, não diagnóstico",
    )
    bullet_card(
        slide,
        Inches(0.5),
        Inches(2.05),
        Inches(6.0),
        Inches(4.7),
        "O cenário",
        [
            "A literatura biomédica cresce mais rápido do que equipes conseguem ler.",
            "Abstracts descrevem a condição do paciente e precisam de um primeiro filtro.",
            "Cinco grupos clínicos recorrem com frequência em papers e prontuários.",
            "Um classificador assistivo padroniza a triagem inicial de texto.",
            "O valor só aparece se o modelo chegar em produção: API, testes e métricas.",
        ],
    )
    bullet_card(
        slide,
        Inches(6.75),
        Inches(2.05),
        Inches(6.05),
        Inches(4.7),
        "Cinco classes-alvo",
        [
            "1  Neoplasms",
            "2  Digestive system diseases",
            "3  Nervous system diseases",
            "4  Cardiovascular diseases",
            "5  General pathological conditions (classe majoritária)",
        ],
    )
    footer(slide, 3)
    notes(
        slide,
        "Situation (0:35–1:15). Enfatize volume de abstracts e que isto não é diagnóstico autônomo.",
    )


def dataset(prs):
    slide = header_slide(
        prs,
        "Situation",
        "Os dados: Medical Abstracts TC Corpus",
        "11.550 treino  ·  2.888 teste  ·  0 nulos  ·  0 duplicados",
    )
    stats = [
        ("11.550", "amostras de treino"),
        ("2.888", "amostras de teste"),
        ("5", "classes clínicas"),
        ("0", "nulos / duplicados"),
    ]
    x = Inches(0.5)
    for n, lab in stats:
        card(slide, x, Inches(2.05), Inches(2.95), Inches(1.55), WHITE)
        add_text_box(slide, x, Inches(2.15), Inches(2.95), Inches(0.85), n, 28, True, TEAL, PP_ALIGN.CENTER)
        add_text_box(
            slide,
            x + Inches(0.1),
            Inches(2.95),
            Inches(2.75),
            Inches(0.5),
            lab,
            13,
            False,
            SLATE,
            PP_ALIGN.CENTER,
        )
        x += Inches(3.15)

    rows = [
        ("Classe", "Treino", "Teste", "Total", "% treino"),
        ("Neoplasms", "2.530", "633", "3.163", "21,9%"),
        ("Digestive system", "1.195", "299", "1.494", "10,3%"),
        ("Nervous system", "1.540", "385", "1.925", "13,3%"),
        ("Cardiovascular", "2.441", "610", "3.051", "21,1%"),
        ("General pathological", "3.844", "961", "4.805", "33,3%"),
    ]
    y = Inches(3.8)
    widths = [Inches(3.6), Inches(1.7), Inches(1.7), Inches(1.7), Inches(1.8)]
    for i, row in enumerate(rows):
        x = Inches(0.5)
        bg = NAVY if i == 0 else (WHITE if i % 2 else RGBColor(0xE8, 0xEE, 0xF4))
        fg = WHITE if i == 0 else SLATE
        for j, cell in enumerate(row):
            rect(slide, x, y, widths[j], Inches(0.42), bg)
            add_text_box(
                slide,
                x + Inches(0.08),
                y + Inches(0.04),
                widths[j] - Inches(0.1),
                Inches(0.35),
                cell,
                12,
                i == 0,
                fg,
            )
            x += widths[j]
        y += Inches(0.42)
    footer(slide, 4)
    notes(
        slide,
        "Dados (1:15–1:40). Destaque o desbalanceamento: classe 5 ~33% e digestive ~10%. "
        "Por isso F1 macro, não só accuracy.",
    )


def task(prs):
    slide = header_slide(
        prs,
        "Task",
        "O que precisava ser entregue",
        "Do CSV bruto até um serviço observável e reproduzível",
    )
    items = [
        ("01", "Explorar", "Volume, colunas, nulos, duplicados e distribuição das classes."),
        ("02", "Treinar", "Baseline NLP clássico, métricas no hold-out oficial e artefato joblib."),
        ("03", "Servir", "API REST com contrato, saúde, confiança e flag de baixa confiança."),
        ("04", "Operar", "Python 3.12, Docker, pytest, Ruff, CI/CD e métricas Prometheus."),
    ]
    y = Inches(2.05)
    for num, title, desc in items:
        card(slide, Inches(0.5), y, Inches(12.3), Inches(1.1), WHITE)
        rect(slide, Inches(0.5), y, Inches(1.15), Inches(1.1), NAVY)
        add_text_box(slide, Inches(0.5), y + Inches(0.28), Inches(1.15), Inches(0.55), num, 18, True, TEAL, PP_ALIGN.CENTER)
        add_text_box(slide, Inches(1.9), y + Inches(0.18), Inches(10.4), Inches(0.4), title, 18, True, NAVY)
        add_text_box(slide, Inches(1.9), y + Inches(0.58), Inches(10.4), Inches(0.4), desc, 14, False, SLATE)
        y += Inches(1.22)
    footer(slide, 5)
    notes(slide, "Task (1:40–2:10). Recorte: MVP ponta a ponta, não o melhor paper de embeddings.")


def architecture(prs):
    slide = header_slide(
        prs,
        "Action",
        "Arquitetura do repositório",
        "Código vivo em src/  ·  dados em data/raw  ·  artefato em models/baseline",
    )
    boxes = [
        ("data/raw", "CSVs train, test\ne labels"),
        ("src/", "loader, TF-IDF,\ntreino e FastAPI"),
        ("models/", "joblib + metrics.json\ndo baseline"),
        ("tests/", "contrato da API\n7 testes"),
        ("CI/CD", "Ruff + pytest\nimagem GHCR"),
    ]
    x = Inches(0.4)
    for title, body in boxes:
        card(slide, x, Inches(2.1), Inches(2.35), Inches(2.05), WHITE)
        rect(slide, x, Inches(2.1), Inches(2.35), Inches(0.5), NAVY)
        add_text_box(slide, x, Inches(2.18), Inches(2.35), Inches(0.4), title, 14, True, WHITE, PP_ALIGN.CENTER)
        add_text_box(
            slide,
            x + Inches(0.1),
            Inches(2.75),
            Inches(2.15),
            Inches(1.2),
            body,
            13,
            False,
            SLATE,
            PP_ALIGN.CENTER,
        )
        x += Inches(2.55)

    bullet_card(
        slide,
        Inches(0.5),
        Inches(4.4),
        Inches(6.0),
        Inches(2.4),
        "Implementado",
        [
            "src/data, features, training e api",
            "scripts/eda.py, README, Makefile, .env.example",
            "Dockerfile 3.12 + docker-compose da API",
            "GitHub Actions: quality gate e push no GHCR",
        ],
    )
    bullet_card(
        slide,
        Inches(6.75),
        Inches(4.4),
        Inches(6.05),
        Inches(2.4),
        "Ainda esqueleto",
        [
            "monitoring/, airflow/ e notebooks/ vazios",
            "Sem testes unitários do treino/loader",
            "Prometheus no endpoint, sem Grafana no Compose",
            "EDA ainda é script de print, não notebook",
        ],
    )
    footer(slide, 6)
    notes(
        slide,
        "Arquitetura (2:10–2:45). Mostre o mapa MLOps e seja honesto: o caminho real é src/ + tests/test_api.py.",
    )


def pipeline(prs):
    slide = header_slide(
        prs,
        "Action",
        "Pipeline de machine learning",
        "Baseline clássico, balanceado, empacotado em um único joblib",
    )
    steps = [
        ("1", "Load", "pandas lê\nmedical_tc_*.csv"),
        ("2", "TF-IDF", "uni+bi, stop EN\nsublinear_tf"),
        ("3", "Classifier", "LogReg balanced\nmax_iter=2000"),
        ("4", "Evaluate", "acc, P/R, F1 macro\nmatriz + JSON"),
        ("5", "Persist", "models.joblib\n+ metrics.json"),
    ]
    x = Inches(0.4)
    for n, title, body in steps:
        card(slide, x, Inches(2.1), Inches(2.35), Inches(2.4), WHITE)
        add_text_box(slide, x, Inches(2.2), Inches(2.35), Inches(0.45), n, 16, True, TEAL, PP_ALIGN.CENTER)
        add_text_box(slide, x, Inches(2.6), Inches(2.35), Inches(0.4), title, 16, True, NAVY, PP_ALIGN.CENTER)
        add_text_box(
            slide,
            x + Inches(0.1),
            Inches(3.1),
            Inches(2.15),
            Inches(1.15),
            body,
            13,
            False,
            SLATE,
            PP_ALIGN.CENTER,
        )
        x += Inches(2.55)

    bullet_card(
        slide,
        Inches(0.5),
        Inches(4.7),
        Inches(12.3),
        Inches(2.1),
        "Decisões técnicas",
        [
            "Pipeline sklearn único: vetorizador e classificador viajam juntos no artefato.",
            "create_vectorizer() em src/features/tfidf.py — ngram (1,2), min_df=2, max_df=0.95.",
            "class_weight='balanced' para compensar a classe 5 (~33%) vs digestive (~10%).",
            "F1 macro como norte; métricas gravadas em models/baseline/metrics.json.",
        ],
    )
    footer(slide, 7)
    notes(
        slide,
        "Modelo (2:45–3:20). Justifique baseline: interpretável, barato e suficiente para o MVP da API.",
    )


def api_slide(prs):
    slide = header_slide(prs, "Action", "API de inferência", "FastAPI + Pydantic + Prometheus")
    endpoints = [
        ("GET /", "Redirect para /docs (Swagger)"),
        ("GET /health", "Liveness e flag model_loaded"),
        ("POST /predict", "Texto → rótulo, nome, confiança, low_confidence"),
        ("GET /metrics", "Contadores, latência e erros"),
    ]
    y = Inches(2.05)
    for ep, desc in endpoints:
        card(slide, Inches(0.5), y, Inches(6.2), Inches(1.05), WHITE)
        add_text_box(slide, Inches(0.75), y + Inches(0.15), Inches(5.7), Inches(0.4), ep, 16, True, NAVY)
        add_text_box(slide, Inches(0.75), y + Inches(0.55), Inches(5.7), Inches(0.4), desc, 13, False, SLATE)
        y += Inches(1.15)

    bullet_card(
        slide,
        Inches(6.95),
        Inches(2.05),
        Inches(5.85),
        Inches(4.65),
        "Contrato e operação",
        [
            "Request: { text } com min_length=1",
            "Response: label, name, confidence, low_confidence",
            "low_confidence = true se confiança < 0,5",
            "Nomes lidos de medical_tc_labels.csv",
            "Validação 422 para payload vazio ou ausente",
            "Histogram de latência em /predict",
            "Fail-fast se o joblib não existir; MODEL_PATH opcional",
        ],
    )
    footer(slide, 8)
    notes(
        slide,
        "API (3:20–3:50). Cite o exemplo de dor no peito: útil para falar de baixa confiança e classe genérica.",
    )


def mlops(prs):
    slide = header_slide(
        prs,
        "Action",
        "Qualidade, container e entrega",
        "Runtime alinhado em Python 3.12 no projeto, CI e Docker",
    )
    cols = [
        (
            "Qualidade",
            [
                "7 testes de contrato da API",
                "Ruff no CI (target py312)",
                "pytest via uv no workflow",
                "Labels conferidos contra o CSV",
            ],
        ),
        (
            "Container",
            [
                "Imagem python:3.12-slim + uv",
                "Copia src/, modelo e labels",
                "CMD uvicorn na porta 8000",
                "Compose com restart e healthcheck",
            ],
        ),
        (
            "CD",
            [
                "Push em main → GHCR",
                "Tag com SHA do commit",
                "packages:write no job",
                "Nome da imagem em lowercase",
            ],
        ),
    ]
    x = Inches(0.5)
    for title, items in cols:
        bullet_card(slide, x, Inches(2.05), Inches(4.0), Inches(4.7), title, items)
        x += Inches(4.15)
    footer(slide, 9)
    notes(
        slide,
        "MLOps (3:50–4:15). Destaque o alinhamento 3.12. O que falta é Grafana/orquestração, não a esteira básica.",
    )


def results(prs):
    slide = header_slide(
        prs,
        "Result",
        "O que já está no ar",
        "Hold-out oficial  ·  último treino gravado em metrics.json",
    )
    kpis = [
        ("0,587", "Accuracy no teste", MUTED),
        ("0,588", "F1 macro", GOLD),
        ("0,629", "Recall macro", TEAL),
        ("7/7", "Testes de contrato", GREEN),
    ]
    x = Inches(0.5)
    for n, lab, color in kpis:
        card(slide, x, Inches(2.05), Inches(2.95), Inches(1.7), WHITE)
        add_text_box(slide, x, Inches(2.15), Inches(2.95), Inches(0.9), n, 28, True, color, PP_ALIGN.CENTER)
        add_text_box(
            slide,
            x + Inches(0.1),
            Inches(3.05),
            Inches(2.75),
            Inches(0.5),
            lab,
            13,
            False,
            SLATE,
            PP_ALIGN.CENTER,
        )
        x += Inches(3.15)

    bullet_card(
        slide,
        Inches(0.5),
        Inches(4.0),
        Inches(6.0),
        Inches(2.8),
        "Entregue",
        [
            "Fluxo E2E: dado → treino → artefato → API",
            "Contrato validado (200, 422, métricas, labels)",
            "Observabilidade básica (Prometheus)",
            "Docker, README, Makefile e CD no GHCR",
        ],
    )
    bullet_card(
        slide,
        Inches(6.75),
        Inches(4.0),
        Inches(6.05),
        Inches(2.8),
        "Limitações honestas",
        [
            "F1 ~0,59 ainda é baixo para triagem clínica real",
            "Classe 5 continua absorvendo casos ambíguos",
            "Confiança pode cair abaixo de 0,5 (flag existe)",
            "Grafana, Airflow e testes de treino ainda faltam",
        ],
    )
    footer(slide, 10)
    notes(
        slide,
        "Resultados (4:15–4:40). Celebre o E2E e não esconda o F1. Credibilidade > marketing.",
    )


def next_steps(prs):
    slide = header_slide(prs, "Result", "Próximos passos", "Do MVP para um serviço mais confiável")
    steps = [
        ("Subir o modelo", "Embeddings biomédicos, SVM linear ou threshold por classe."),
        ("Fechar observabilidade", "Prometheus + Grafana no Compose; alertas de latência e erro."),
        ("Testes de ML", "Unitários do loader/vectorizer e regressão de métricas no CI."),
        ("Orquestração", "Job de retreino reproduzível (hoje o treino é script local)."),
        ("Governança", "Disclaimer visível: apoio à triagem, não diagnóstico. Auditoria de erros."),
    ]
    y = Inches(2.05)
    for i, (title, desc) in enumerate(steps, 1):
        card(slide, Inches(0.5), y, Inches(12.3), Inches(0.9), WHITE)
        rect(slide, Inches(0.5), y, Inches(0.9), Inches(0.9), NAVY)
        add_text_box(
            slide,
            Inches(0.5),
            y + Inches(0.22),
            Inches(0.9),
            Inches(0.5),
            f"{i:02d}",
            16,
            True,
            TEAL,
            PP_ALIGN.CENTER,
        )
        add_text_box(slide, Inches(1.65), y + Inches(0.12), Inches(10.8), Inches(0.35), title, 16, True, NAVY)
        add_text_box(slide, Inches(1.65), y + Inches(0.48), Inches(10.8), Inches(0.35), desc, 13, False, SLATE)
        y += Inches(0.98)
    footer(slide, 11)
    notes(slide, "Próximos passos (4:40–4:55). Se faltar tempo, fale só modelo + Grafana.")


def close(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, NAVY)
    rect(slide, 0, 0, Inches(0.18), H, TEAL)
    add_text_box(slide, Inches(0.7), Inches(1.8), Inches(12), Inches(0.4), "RESULT  ·  MENSAGEM FINAL", 14, True, TEAL)
    add_text_box(
        slide,
        Inches(0.7),
        Inches(2.3),
        Inches(12),
        Inches(1.4),
        "Do corpus ao endpoint\nem um único fluxo.",
        36,
        True,
        WHITE,
    )
    add_text_box(
        slide,
        Inches(0.7),
        Inches(4.2),
        Inches(11.5),
        Inches(1.4),
        "O Med-Triage classifica abstracts, devolve confiança e nasce com testes,\n"
        "container e CD. O próximo salto é qualidade do modelo e observabilidade visual.",
        18,
        False,
        RGBColor(0xC9, 0xD6, 0xE4),
    )
    add_text_box(slide, Inches(0.7), Inches(6.3), Inches(12), Inches(0.4), "Obrigado  ·  perguntas", 20, True, GOLD)
    notes(slide, "Fechamento (4:55–5:00). Convide perguntas sobre modelo, API ou CI.")


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    cover(prs)
    agenda(prs)
    situation(prs)
    dataset(prs)
    task(prs)
    architecture(prs)
    pipeline(prs)
    api_slide(prs)
    mlops(prs)
    results(prs)
    next_steps(prs)
    close(prs)
    prs.save(OUT)
    print(f"Salvo em {OUT}")


if __name__ == "__main__":
    main()
